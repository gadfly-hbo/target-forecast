# -*- coding: utf-8 -*-
"""数据导入与质量检查（提案 §3.3–3.5）：手工分桶 / CSV / XLSX → DataSnapshot。

金额一律分（整数）。质量检查按 error/warning 分级；error 不阻断快照生成，
但 snapshot_to_baseline / load_spec 会把 error 带给上层决定是否继续。
"""
from __future__ import annotations

import csv
import json
import os
from dataclasses import dataclass, field

from .spec import Basket


@dataclass
class DataSnapshot:
    """基准分布的不可变快照及口径信息。"""
    snapshot_id: str
    source: str
    caliber: dict = field(default_factory=dict)
    baskets: list = field(default_factory=list)
    issues: list = field(default_factory=list)
    row_count: int = 0


# ---------------- 读取器 ----------------

def read_csv_file(path: str) -> tuple[list[str], list[dict]]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        reader = csv.DictReader(fh)
        headers = list(reader.fieldnames or [])
        rows = [dict(r) for r in reader]
    return headers, rows


def read_xlsx_file(path: str, sheet: str | None = None) -> tuple[list[str], list[dict]]:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    except StopIteration:
        wb.close()
        return [], []
    out = []
    for values in rows_iter:
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append({h: v for h, v in zip(headers, values) if h})
    wb.close()
    return headers, out


# ---------------- 工具 ----------------

def _issue(level, code, message):
    return {"level": level, "code": code, "message": message}


def _int_cents_or_issue(value, issues, code_non_int="non_integer_cents", label="金额"):
    """把单元格值转整数分；非整数分 → error；明显是“元”量级 → 疑似单位混杂警告。"""
    try:
        f = float(value)
    except (TypeError, ValueError):
        issues.append(_issue("error", code_non_int, f"{label}不是数字：{value!r}"))
        return None
    if abs(f - round(f)) > 1e-9:
        issues.append(_issue("error", code_non_int, f"{label}必须是整数分，得到 {value}"))
        return None
    cents = int(round(f))
    if 0 < cents < 100:
        issues.append(_issue("warning", "currency_unit_suspect",
                             f"{label}={cents} 分（<1 元）：疑似把“元”填进了“分”列"))
    return cents


def _num(value, default=None):
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


# ---------------- 分桶导入 ----------------

def _rows_arg(data):
    """兼容 (headers, rows) 元组或 rows 列表两种入参。"""
    return data[1] if isinstance(data, tuple) else data


def import_buckets(data, meta: dict) -> DataSnapshot:
    """分桶表（bucket_min/max 可空、avg_amount_cents 必填、orders 或 weight 二选一）。"""
    rows = _rows_arg(data)
    issues: list = []
    buckets: list[dict] = []
    weight_mode = bool(rows) and all(
        (r.get("orders") in (None, "")) and r.get("weight") not in (None, "") for r in rows)

    for idx, r in enumerate(rows, start=2):
        avg = _int_cents_or_issue(r.get("avg_amount_cents"), issues,
                                  label=f"第{idx}行桶均金额")
        if avg is None:
            continue
        if avg <= 0:
            issues.append(_issue("error", "nonpositive_amount", f"第{idx}行桶均金额 ≤ 0"))
            continue
        bmin = _int_cents_or_issue(r["bucket_min_cents"], issues,
                                   label=f"第{idx}行桶下限") if r.get("bucket_min_cents") not in (None, "") else None
        bmax = _int_cents_or_issue(r["bucket_max_cents"], issues,
                                   label=f"第{idx}行桶上限") if r.get("bucket_max_cents") not in (None, "") else None
        margin = _num(r.get("avg_margin_rate"))
        if margin is not None and not 0.0 <= margin <= 1.0:
            issues.append(_issue("error", "illegal_probability", f"第{idx}行毛利率必须在 [0,1]"))
            margin = None
        if weight_mode:
            w = _num(r.get("weight"))
            if w is None or w < 0:
                issues.append(_issue("error", "illegal_weight", f"第{idx}行权重非法"))
                continue
        else:
            orders = _num(r.get("orders"))
            if orders is None or orders < 0:
                issues.append(_issue("error", "missing_value", f"第{idx}行缺少订单数/权重"))
                continue
            w = orders
        buckets.append({"avg": avg, "w": w, "margin": margin, "bmin": bmin, "bmax": bmax})

    total_w = sum(b["w"] for b in buckets)
    if not buckets:
        issues.append(_issue("error", "empty_input", "没有可用的分桶行"))
    elif weight_mode and abs(total_w - 1.0) > 1e-6:
        issues.append(_issue("error", "weight_not_conserved",
                             f"权重合计必须为 1，实际 {total_w:.6f}"))
    elif not weight_mode and total_w <= 0:
        issues.append(_issue("error", "weight_not_conserved", "订单数合计必须 > 0"))

    baskets = []
    for b in buckets:
        norm = b["w"] / total_w if total_w > 0 else 0.0
        baskets.append(Basket(
            amount_cents=b["avg"], weight=norm, margin_rate=b["margin"],
            label=f"桶[{b['bmin'] if b['bmin'] is not None else '−∞'},"
                  f"{b['bmax'] if b['bmax'] is not None else '+∞'}]均值{b['avg']}",
            bucket_min_cents=b["bmin"], bucket_max_cents=b["bmax"],
        ))

    return DataSnapshot(
        snapshot_id=str(meta.get("snapshot_id", "unnamed-snapshot")),
        source=str(meta.get("source", "manual")),
        caliber=dict(meta.get("caliber", {})),
        baskets=baskets, issues=issues, row_count=len(rows),
    )


# ---------------- 订单明细导入 ----------------

def import_orders(data, meta: dict) -> DataSnapshot:
    """订单明细（order_id、amount_cents 券前商品金额、margin_cents 可选、refunded 可选）→ 逐单聚合。"""
    rows = _rows_arg(data)
    issues: list = []
    seen_ids: set = set()
    by_amount: dict[int, dict] = {}
    refunded_seen = False
    count = 0

    for idx, r in enumerate(rows, start=2):
        oid = r.get("order_id")
        if oid in (None, ""):
            issues.append(_issue("error", "missing_value", f"第{idx}行缺少订单号"))
            continue
        if oid in seen_ids:
            issues.append(_issue("error", "duplicate_order", f"订单号重复：{oid}"))
            continue
        seen_ids.add(oid)
        amount = _int_cents_or_issue(r.get("amount_cents"), issues,
                                     label=f"订单{oid}金额")
        if amount is None:
            continue
        if amount <= 0:
            issues.append(_issue("error", "nonpositive_amount", f"订单{oid}券前金额 ≤ 0"))
            continue
        margin = _int_cents_or_issue(r.get("margin_cents"), issues,
                                     label=f"订单{oid}毛利") if r.get("margin_cents") not in (None, "") else None
        if r.get("refunded") not in (None, "", 0, "0", False):
            refunded_seen = True
        slot = by_amount.setdefault(amount, {"n": 0, "margin_sum": 0.0, "margin_n": 0})
        slot["n"] += 1
        if margin is not None:
            slot["margin_sum"] += margin / amount
            slot["margin_n"] += 1
        count += 1

    if refunded_seen and not meta.get("refund_handling"):
        issues.append(_issue("warning", "refund_unspecified",
                             "存在退款单但未说明退款处理口径（含/剔除/净额），请先确认"))

    total = sum(s["n"] for s in by_amount.values())
    baskets = []
    for amount in sorted(by_amount):
        s = by_amount[amount]
        rate = (s["margin_sum"] / s["margin_n"]) if s["margin_n"] else None
        baskets.append(Basket(amount_cents=amount, weight=s["n"] / total,
                              margin_rate=rate, label=f"A={amount}×{s['n']}"))

    if not baskets:
        issues.append(_issue("error", "empty_input", "没有可用订单"))

    return DataSnapshot(
        snapshot_id=str(meta.get("snapshot_id", "unnamed-snapshot")),
        source=str(meta.get("source", "orders")),
        caliber=dict(meta.get("caliber", {})),
        baskets=baskets, issues=issues, row_count=len(rows),
    )


# ---------------- 快照 → 基准 ----------------

def snapshot_to_baseline(snapshot: DataSnapshot, visitors: int | None,
                         conversion_rate: float) -> dict:
    """快照转 ScenarioSpec.baseline 的 dict 片段；error 未决时拒绝。"""
    errors = [i for i in snapshot.issues if i["level"] == "error"]
    if errors:
        raise ValueError("数据存在 error 级问题，先修复再测算：" + "；".join(e["message"] for e in errors))
    return {
        "visitors": visitors,
        "conversion_rate": conversion_rate,
        "data_snapshot_id": snapshot.snapshot_id,
        "baskets": [{
            "amount_cents": b.amount_cents, "weight": b.weight,
            "margin_rate": b.margin_rate, "label": b.label,
            "bucket_min_cents": b.bucket_min_cents, "bucket_max_cents": b.bucket_max_cents,
        } for b in snapshot.baskets],
    }


# ---------------- 口径交叉检查 ----------------

def check_cost_overlap(margin_includes_fulfillment: bool, fixed_fee_cents: float,
                       fee_rate: float) -> bool:
    """M16：已含履约的毛利 + 履约费 = 成本重复扣减。"""
    return bool(margin_includes_fulfillment and (fixed_fee_cents > 0 or fee_rate > 0))


def check_reach_overlap(audience_already_reached: bool, reach_base: float) -> bool:
    """触达率重复计算：N 已限定为实际触达者，却又乘 e<1。"""
    return bool(audience_already_reached and reach_base < 1.0)


def check_window_consistency(orders_window: str | None, traffic_window: str | None) -> bool:
    return bool(orders_window and traffic_window and orders_window != traffic_window)


# ---------------- 模板生成 ----------------

_BUCKET_TEMPLATE_ROWS = [
    [10000, 14999, 12000, 350, 0.30],
    [15000, 19999, 17000, 400, 0.28],
    [None, None, 9900, 250, 0.30],
]
_ORDER_TEMPLATE_ROWS = [
    ["o-0001", 9900, 3000, 0],
    ["o-0002", 15800, 4700, 0],
    ["o-0003", 12000, 3500, 1],
]


def generate_templates(out_dir: str) -> list[str]:
    """生成分桶与订单明细两种 XLSX 模板（含说明页与示例行）。"""
    from openpyxl import Workbook
    os.makedirs(out_dir, exist_ok=True)
    written = []

    wb = Workbook()
    ws = wb.active
    ws.title = "分桶数据"
    ws.append(BUCKET_HEADER := ["bucket_min_cents", "bucket_max_cents", "avg_amount_cents",
                                "orders", "avg_margin_rate", "weight"])
    for row in _BUCKET_TEMPLATE_ROWS:
        ws.append(row + [None])
    note = wb.create_sheet("填数说明")
    note.append(["列", "说明"])
    for col, desc in [
        ("bucket_min_cents", "桶下限（分，可空=开区间）"),
        ("bucket_max_cents", "桶上限（分，可空）"),
        ("avg_amount_cents", "桶内平均券前可用商品金额（分，必填）——不是实付金额"),
        ("orders", "桶内订单数（与 weight 二选一）"),
        ("avg_margin_rate", "桶均商品毛利率 [0,1]（可空）"),
        ("weight", "直接给权重时填写（与 orders 二选一，合计必须=1）"),
    ]:
        note.append([col, desc])
    path = os.path.join(out_dir, "优惠券测算_分桶模板.xlsx")
    wb.save(path)
    written.append(os.path.basename(path))

    wb = Workbook()
    ws = wb.active
    ws.title = "订单明细"
    ws.append(["order_id", "amount_cents", "margin_cents", "refunded"])
    for row in _ORDER_TEMPLATE_ROWS:
        ws.append(row)
    note = wb.create_sheet("填数说明")
    note.append(["列", "说明"])
    for col, desc in [
        ("order_id", "订单号（用于查重，不进入测算）"),
        ("amount_cents", "本次券前可用商品金额（分，必填），不含运费与其他优惠"),
        ("margin_cents", "商品毛利（分，可空；空则用场景级毛利率）"),
        ("refunded", "是否退款单（0/1；存在 1 时需在导入确认退款口径）"),
    ]:
        note.append([col, desc])
    path = os.path.join(out_dir, "优惠券测算_订单明细模板.xlsx")
    wb.save(path)
    written.append(os.path.basename(path))
    return written
